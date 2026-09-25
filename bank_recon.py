#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""銀行匯款對帳 — 把官網「匯款訂單」CSV 整理進銀行匯款對帳總表。

用法:
    python bank_recon.py <匯款訂單.csv> [--total 對帳總表.xlsx] [--dept 電圖部]
    python bank_recon.py <匯款訂單.csv> --lookup "客戶訂單+代號_銷貨明細.xlsx"
    python bank_recon.py --lookup "客戶訂單+代號_銷貨明細.xlsx"   # 只補客戶代號
"""
import argparse
import csv
import os
import re
import sys
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment

SHEET_NAME = '銀行匯款對帳'
CORE_HEADERS = ['交易日', '存款金額', '摘要', '備註', '部門', '網站訂單編號', '賣家備註']
CUST_HEADER = '客戶代號'
ENCODINGS = ['cp950', 'big5', 'utf-8-sig', 'utf-8']
DEFAULT_TOTAL = '銀行匯款對帳資料-總表.xlsx'


# ---------------------------------------------------------------- CSV 讀取
def read_csv(path):
    last = None
    for enc in ENCODINGS:
        try:
            with open(path, encoding=enc) as f:
                return list(csv.reader(f)), enc
        except UnicodeDecodeError as e:
            last = e
    raise last


def build_index(header):
    idx = {}
    for i, h in enumerate(header):
        h = h.strip()
        if h and h not in idx:
            idx[h] = i
    return idx


def g(row, idx, name):
    i = idx.get(name)
    if i is None or i >= len(row):
        return ''
    return row[i].strip()


def clean_amount(s):
    s = (s or '').replace(',', '').strip()
    if not s:
        return None
    try:
        v = float(s)
    except ValueError:
        return s
    return int(v) if v == int(v) else v


def unwrap(s):
    if s.startswith('="') and s.endswith('"'):
        return s[2:-1]
    return s


def parse_date(s):
    d = (s or '').split(' ')[0].replace('/', '-')
    if not d:
        return ''
    y, m, day = [int(x) for x in d.split('-')]
    return datetime(y, m, day).date()


def correct_last5(note):
    """從賣家備註抓「正確後5碼」的 5 位數字，找不到回 None。"""
    if '正確' not in note:
        return None
    m = re.search(r'\d{5}', note)
    return m.group(0) if m else None


def parse_records(rows, dept):
    idx = build_index(rows[0])
    records = []
    for row in rows[1:]:
        if not any(c.strip() for c in row):
            continue
        order_no = g(row, idx, '訂單編號')
        if not order_no:
            continue
        seller_note = g(row, idx, '賣家備註')
        summary = unwrap(g(row, idx, '轉帳後五碼'))
        fixed = correct_last5(seller_note)
        if fixed:
            summary = fixed
        records.append({
            '交易日': parse_date(g(row, idx, '匯款日期')),
            '存款金額': clean_amount(g(row, idx, '匯款金額')),
            '摘要': summary,
            '備註': g(row, idx, '匯款銀行'),
            '部門': dept,
            '網站訂單編號': order_no,
            '賣家備註': seller_note,
        })
    return records


# ---------------------------------------------------------------- 總表存取
def load_or_create(path):
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.create_sheet(SHEET_NAME)
        if ws.max_row == 0 or ws.cell(row=1, column=1).value is None:
            ws.append(CORE_HEADERS)
        return wb, ws
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(CORE_HEADERS)
    return wb, ws


def col_map(ws):
    m = {}
    for i, c in enumerate(ws[1], 1):
        if c.value not in (None, ''):
            m[str(c.value).strip()] = i
    return m


def ensure_column(ws, name, after_name):
    """確保總表有 name 欄；沒有就在 after_name 後面插一欄，回傳欄序號。"""
    m = col_map(ws)
    if name in m:
        return m[name]
    after = m.get(after_name)
    pos = (after + 1) if after else ws.max_column + 1
    ws.insert_cols(pos)
    ws.cell(row=1, column=pos, value=name)
    return pos


def style_rows(ws, start, cm):
    for i in range(start, ws.max_row + 1):
        if '交易日' in cm:
            c = ws.cell(row=i, column=cm['交易日'])
            c.alignment = Alignment(horizontal='center')
            if c.value not in (None, ''):
                c.number_format = 'yyyy/mm/dd'
        if '存款金額' in cm:
            c = ws.cell(row=i, column=cm['存款金額'])
            c.alignment = Alignment(horizontal='right')
            if isinstance(c.value, (int, float)):
                c.number_format = '#,##0'
        if '摘要' in cm:
            c = ws.cell(row=i, column=cm['摘要'])
            if c.value not in (None, ''):
                c.value = str(c.value)
                c.number_format = '@'
                c.alignment = Alignment(horizontal='center')


# ---------------------------------------------------------------- 客戶代號
def build_lookup(path):
    """讀『客戶訂單 | 客戶代號』xlsx，回傳 {客戶訂單: 客戶代號} 與衝突數。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    look = {}
    conflicts = 0
    for order, code in ws.iter_rows(min_row=2, values_only=True):
        if order is None or code is None:
            continue
        o = str(order).strip()
        c = str(code).strip()
        if not o:
            continue
        if o in look and look[o] != c:
            conflicts += 1
        look[o] = c
    wb.close()
    return look, conflicts


def apply_lookup(ws, lookup_path):
    look, conflicts = build_lookup(lookup_path)
    pos = ensure_column(ws, CUST_HEADER, '網站訂單編號')
    cm = col_map(ws)
    order_pos = cm.get('網站訂單編號')
    if order_pos is None:
        raise SystemExit('總表找不到「網站訂單編號」欄，無法比對客戶代號')

    matched, missed = 0, []
    for i in range(2, ws.max_row + 1):
        v = ws.cell(row=i, column=order_pos).value
        o = str(v).strip() if v is not None else ''
        code = look.get(o, '')
        ws.cell(row=i, column=pos, value=code if code else None)
        if code:
            matched += 1
        elif o:
            missed.append((i, o))

    ws.column_dimensions[openpyxl.utils.get_column_letter(pos)].width = 14
    return {
        'lookup': len(look),
        'conflicts': conflicts,
        'matched': matched,
        'total': ws.max_row - 1,
        'missed': missed,
    }


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description='銀行匯款對帳：匯款訂單 CSV → 對帳總表')
    ap.add_argument('csv', nargs='?', help='官網匯出的匯款訂單 CSV')
    ap.add_argument('--total', default=DEFAULT_TOTAL, help='對帳總表檔名')
    ap.add_argument('--dept', default='電圖部', help='部門固定值')
    ap.add_argument('--overwrite', action='store_true', help='清空總表既有資料列後重寫（保留表頭）')
    ap.add_argument('--lookup', help='「客戶訂單｜客戶代號」xlsx，用來比對回填客戶代號')
    args = ap.parse_args()

    if not args.csv and not args.lookup:
        ap.error('至少要給 CSV，或加 --lookup 補客戶代號')

    try:
        wb, ws = load_or_create(args.total)

        if args.csv:
            rows, enc = read_csv(args.csv)
            if not rows:
                print('CSV 是空的'); return 1
            records = parse_records(rows, args.dept)
            if not records:
                print('沒有可輸出的資料'); return 1

            if args.overwrite and ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)

            if args.lookup:
                ensure_column(ws, CUST_HEADER, '網站訂單編號')
            cm = col_map(ws)
            start = ws.max_row + 1
            for rec in records:
                r = ws.max_row + 1
                for name, val in rec.items():
                    if name in cm:
                        ws.cell(row=r, column=cm[name], value=val)
            style_rows(ws, start, cm)

            fixed = sum(1 for r in records if correct_last5(r['賣家備註']))
            print('來源: %s（編碼 %s）' % (os.path.basename(args.csv), enc))
            print('寫入 %d 筆 → %s（第 %d~%d 列）' % (len(records), args.total, start, ws.max_row))
            if fixed:
                print('摘要以「正確後5碼」覆蓋: %d 筆' % fixed)

        if args.lookup:
            rep = apply_lookup(ws, args.lookup)
            print('客戶代號比對: 查找表 %d 筆（衝突 %d）→ 帶入 %d/%d 列'
                  % (rep['lookup'], rep['conflicts'], rep['matched'], rep['total']))
            if rep['missed']:
                print('  未比對到: %s' % rep['missed'])

        wb.save(args.total)
    except PermissionError:
        print('寫入失敗：%s 正被 Excel 開啟，請先關閉再執行。' % args.total)
        return 1

    return 0


if __name__ == '__main__':
    sys.exit(main())
